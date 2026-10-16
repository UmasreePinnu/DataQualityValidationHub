import time
import streamlit as st
import pandas as pd
import numpy as np
from io import BytesIO
import openpyxl
from openpyxl.styles import PatternFill, Font
from metrics_tracker import log_event


def run():
    st.title("📁 File Comparator")
    st.caption("Upload paired _Source.xlsx and _Target.xlsx files to compare differences")

    st.subheader("Upload files")
    uploaded_files = st.file_uploader(
        "Drop Excel files here",
        type=["xlsx"],
        accept_multiple_files=True,
        label_visibility="collapsed",
    )

    if "comparator_results" not in st.session_state:
        st.session_state["comparator_results"] = {}

    if not uploaded_files:
        st.info("Upload paired files: `TableName_Source.xlsx` and `TableName_Target.xlsx`", icon="📂")
        return

    file_dict = {f.name: f for f in uploaded_files}

    all_bases = set()
    for f in file_dict.keys():
        base = f.replace("_Source.xlsx", "").replace("_Target.xlsx", "")
        all_bases.add(base)

    missing_any = False
    for base in sorted(all_bases):
        src_exists = f"{base}_Source.xlsx" in file_dict
        trg_exists = f"{base}_Target.xlsx" in file_dict

        if not src_exists:
            st.error(f"Missing **Source** file for: {base}", icon="🚨")
            missing_any = True
        if not trg_exists:
            st.error(f"Missing **Target** file for: {base}", icon="🚨")
            missing_any = True

    if missing_any:
        st.stop()

    if st.button("▶ Run Comparison", type="primary"):
        t_all_start = time.perf_counter()
        yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        bold_font = Font(bold=True)
        progress = st.progress(0, text="Comparing files…")
        results = {}
        per_table_timings = []

        for idx, base in enumerate(sorted(all_bases)):
            t_base_start = time.perf_counter()
            src_name = f"{base}_Source.xlsx"
            trg_name = f"{base}_Target.xlsx"

            df1 = pd.read_excel(file_dict[src_name], dtype=str)
            df2 = pd.read_excel(file_dict[trg_name], dtype=str)

            df1.columns = df1.columns.str.upper().str.strip()
            df2.columns = df2.columns.str.upper().str.strip()

            KEY = "KEY" if "KEY" in df1.columns else df1.columns[0]

            src_dups = df1[KEY].duplicated().sum()
            trg_dups = df2[KEY].duplicated().sum()

            if src_dups > 0 or trg_dups > 0:
                st.error(
                    f"⛔ **Comparison Aborted for {base}:** Duplicate keys discovered! "
                    f"Source file contains `{src_dups}` duplicate keys. "
                    f"Target file contains `{trg_dups}` duplicate keys. Please fix keys to proceed.",
                    icon="🚨"
                )
                st.stop()

            src_only_cols = list(set(df1.columns) - set(df2.columns))
            trg_only_cols = list(set(df2.columns) - set(df1.columns))
            common_cols = [c for c in df1.columns if c in df2.columns]

            for df in [df1, df2]:
                for col in common_cols:
                    df[col] = df[col].fillna("NULL").astype(str).str.strip()
                    df.loc[df[col] == "", col] = ""

            src_keys = set(df1[KEY])
            trg_keys = set(df2[KEY])

            miss_trg_keys = list(src_keys - trg_keys)
            miss_src_keys = list(trg_keys - src_keys)

            cols_reordered = [KEY] + [c for c in common_cols if c != KEY]
            df_miss_trg = df1[df1[KEY].isin(miss_trg_keys)][cols_reordered]
            df_miss_src = df2[df2[KEY].isin(miss_src_keys)][cols_reordered]

            df1c = df1[~df1[KEY].isin(miss_trg_keys)]
            df2c = df2[~df2[KEY].isin(miss_src_keys)]
            merged = pd.merge(df1c, df2c, on=KEY, how="inner", suffixes=("_SRC", "_TRG"))

            df_all = pd.DataFrame()
            df_all[KEY] = merged[KEY]

            detailed_diff_rows = []

            for col in [c for c in common_cols if c != KEY]:
                sc, tc = col + "_SRC", col + "_TRG"
                mask = merged[sc] != merged[tc]
                df_all[col] = np.where(mask, merged[sc] + " → " + merged[tc], merged[sc])

                if mask.any():
                    mismatched_sub = merged[mask]
                    for _, row in mismatched_sub.iterrows():
                        val_str = f"{row[sc]} → {row[tc]}"
                        detailed_diff_rows.append({
                            "Column Name": col,
                            "Issue": val_str,
                            "KEY Value": row[KEY],
                        })

            # Build detailed report dataframe with single representative row per Column Name
            if detailed_diff_rows:
                df_temp = pd.DataFrame(detailed_diff_rows)
                counts = df_temp.groupby("Column Name").size().reset_index(name="Issue Count")
                df_temp_unique = df_temp.drop_duplicates(subset=["Column Name"]).copy()
                df_diff_cols_report = pd.merge(df_temp_unique, counts, on="Column Name", how="left")
                df_diff_cols_report = df_diff_cols_report[
                    ["Column Name", "Issue", "KEY Value", "Issue Count"]
                ]
            else:
                df_diff_cols_report = pd.DataFrame(
                    columns=["Column Name", "Issue", "KEY Value", "Issue Count"]
                )

            row_has_issue_mask = df_all.apply(lambda r: any("→" in str(v) for v in r), axis=1)
            df_filtered = df_all[row_has_issue_mask].copy()

            output = BytesIO()
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                # Sheet 1: Missing Columns
                pd.DataFrame({"Missing in Target": src_only_cols}).to_excel(writer, sheet_name="Missing_Columns",
                                                                            index=False, startrow=1)
                pd.DataFrame({"Missing in Source": trg_only_cols}).to_excel(writer, sheet_name="Missing_Columns",
                                                                            index=False,
                                                                            startrow=len(src_only_cols) + 5)

                # Sheet 2: Missing Records
                wb = writer.book
                ws_miss = wb.create_sheet(title="Missing_Records")
                ws_miss.append(["Missing in Target"])
                ws_miss.cell(row=ws_miss.max_row, column=1).font = bold_font
                ws_miss.append(list(df_miss_trg.columns))
                for cell in ws_miss[ws_miss.max_row]:
                    cell.font = bold_font
                for row in df_miss_trg.values.tolist():
                    ws_miss.append(row)

                ws_miss.append([])
                ws_miss.append([])
                ws_miss.append(["Missing in Source"])
                ws_miss.cell(row=ws_miss.max_row, column=1).font = bold_font
                ws_miss.append(list(df_miss_src.columns))
                for cell in ws_miss[ws_miss.max_row]:
                    cell.font = bold_font
                for row in df_miss_src.values.tolist():
                    ws_miss.append(row)

                # Sheet 3: All Data
                df_all.to_excel(writer, sheet_name="All_Data", index=False)

                # Sheet 4: Differences Columns
                df_diff_cols_report.to_excel(writer, sheet_name="Differences_Columns", index=False)

                # Sheet 5: Filtered Differences
                df_filtered.to_excel(writer, sheet_name="Filtered_Differences", index=False)

                def _apply_arrow_highlights(ws, df):
                    if df.empty:
                        return
                    mask = df.apply(lambda col: col.astype(str).str.contains("→", na=False)).values
                    row_idxs, col_idxs = mask.nonzero()
                    for r, c in zip(row_idxs, col_idxs):
                        ws.cell(row=int(r) + 2, column=int(c) + 1).fill = yellow_fill

                if "All_Data" in wb.sheetnames:
                    _apply_arrow_highlights(wb["All_Data"], df_all)
                if "Filtered_Differences" in wb.sheetnames:
                    _apply_arrow_highlights(wb["Filtered_Differences"], df_filtered)

                # Highlight the "Issue" column cells in Differences_Columns sheet (Column B)
                if "Differences_Columns" in wb.sheetnames and not df_diff_cols_report.empty:
                    ws_diff = wb["Differences_Columns"]
                    for r in range(2, ws_diff.max_row + 1):
                        ws_diff.cell(row=r, column=2).fill = yellow_fill

            base_elapsed = round(time.perf_counter() - t_base_start, 2)
            per_table_timings.append({"table": base, "duration_sec": base_elapsed})

            results[base] = {
                "data": output.getvalue(),
                "miss_src_cols": len(trg_only_cols),
                "miss_trg_cols": len(src_only_cols),
                "miss_src_rows": len(miss_src_keys),
                "miss_trg_rows": len(miss_trg_keys),
                "diff_rows": len(df_filtered),
            }
            progress.progress((idx + 1) / len(all_bases))

        st.session_state["comparator_results"] = results
        progress.empty()

        total_elapsed = time.perf_counter() - t_all_start
        log_event(
            feature="File Comparator",
            duration_sec=total_elapsed,
            details={
                "output_files_count": len(results),
                "per_table_timings": per_table_timings,
            },
        )

    if st.session_state["comparator_results"]:
        st.divider()
        st.subheader("Results")

        summary_rows = []
        for name, r in st.session_state["comparator_results"].items():
            total_issues = (r["miss_src_cols"] + r["miss_trg_cols"] +
                            r["miss_src_rows"] + r["miss_trg_rows"] + r["diff_rows"])
            status = "✅ Match" if total_issues == 0 else "❌ Issues found"

            summary_rows.append({
                "File pair": name,
                "Missing in SRC (Cols)": r["miss_src_cols"],
                "Missing in TRG (Cols)": r["miss_trg_cols"],
                "Missing in SRC (Rows)": r["miss_src_rows"],
                "Missing in TRG (Rows)": r["miss_trg_rows"],
                "Diff rows": r["diff_rows"],
                "Status": status,
            })

        st.dataframe(pd.DataFrame(summary_rows), use_container_width=True, hide_index=True)

        st.markdown("**Download reports**")
        item_list = list(st.session_state["comparator_results"].items())
        cols = st.columns(3)
        for i, (name, data_dict) in enumerate(item_list):
            with cols[i % 3]:
                st.download_button(
                    f"⬇ {name}",
                    data=data_dict["data"],
                    file_name=f"{name}_Comparison.xlsx",
                    key=f"dl_btn_{name}"
                )


if __name__ == "__main__":
    run()