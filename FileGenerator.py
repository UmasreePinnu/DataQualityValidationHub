import re
import time
from io import BytesIO
import openpyxl
import pandas as pd
import pyodbc
import streamlit as st
from conn_picker import render_picker
from db_utils import engine_from_profile
from metrics_tracker import log_event

_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _sanitize(value):
    if isinstance(value, str):
        return _ILLEGAL_CHARS_RE.sub("", value)
    return value


def _build_excel_file(sheet_name, df):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    for col_idx, col_name in enumerate(df.columns, 1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row_values in enumerate(df.values, 2):
        for col_idx, value in enumerate(row_values, 1):
            if pd.isna(value):
                cell_value = "NULL"
            elif str(value).strip() == "":
                cell_value = None
            else:
                cell_value = _sanitize(value)
            ws.cell(row=row_idx, column=col_idx, value=cell_value)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _decode_bytes_latin1(value):
    if value is None:
        return None
    if isinstance(value, bytes):
        return value.decode("latin1", errors="replace")
    return value


def safe_execute_query(engine, query_str: str) -> pd.DataFrame:
    raw_conn = engine.raw_connection()
    try:
        dbapi_conn = getattr(raw_conn, "driver_connection", raw_conn)

        if hasattr(dbapi_conn, "add_output_converter"):
            for sql_type in (pyodbc.SQL_CHAR, 12, -1):
                dbapi_conn.add_output_converter(sql_type, _decode_bytes_latin1)

        cursor = dbapi_conn.cursor()
        cursor.execute(query_str)

        if not cursor.description:
            return pd.DataFrame()

        columns = [col[0] for col in cursor.description]
        rows = cursor.fetchall()

        data = []
        for row in rows:
            clean_row = []
            for val in row:
                if isinstance(val, bytes):
                    val = val.decode("latin1", errors="replace")
                clean_row.append(val)
            data.append(clean_row)

        return pd.DataFrame(data, columns=columns)
    finally:
        raw_conn.close()


def run():
    st.title("📄 File Generator")
    st.caption("Run SQL queries against a database and export results to Excel")

    if "generated_results" not in st.session_state:
        st.session_state["generated_results"] = None

    st.subheader("Connection")
    profile = render_picker("Database connection", key="file_gen_conn")

    st.divider()

    st.subheader("Output format")
    st.caption("Store results in:")
    output_mode = st.radio(
        "output_mode_label",
        ["One Excel file (multiple sheets)", "Separate Excel files per query"],
        label_visibility="collapsed",
        index=0,
    )

    st.divider()

    st.subheader("SQL queries")
    st.caption(
        "Separate queries with `;`. Add a `-- Sheet name` comment above each query to name its sheet"
    )

    query_text = st.text_area(
        "sql_input",
        placeholder="-- SHEET_NAME\nSELECT * FROM table;",
        height=300,
        label_visibility="collapsed",
    )

    if st.button("▶ Run Queries", type="primary"):
        if not profile:
            st.error("Please select a connection profile.")
            return
        if not query_text.strip():
            st.error("Please enter at least one SQL query.")
            return

        t_start = time.perf_counter()
        try:
            raw_queries = [
                q.strip() for q in query_text.split(";") if q.strip()
            ]
            results = []
            items_summary = []

            with st.spinner("Executing queries... Please check for MFA prompt if needed."):
                engine = engine_from_profile(profile)

                for q in raw_queries:
                    match = re.search(r"--\s*([\w\d_]+)", q)
                    sheet_name = (
                        match.group(1)
                        if match
                        else f"Sheet_{len(results) + 1}"
                    )

                    t_query_start = time.perf_counter()
                    df = safe_execute_query(engine, q)
                    query_duration = round(time.perf_counter() - t_query_start, 2)

                    if not df.empty:
                        results.append((sheet_name, df))
                        items_summary.append({
                            "name": sheet_name,
                            "rows": len(df),
                            "columns": len(df.columns),
                            "duration_sec": query_duration,
                        })
                    else:
                        st.info(f"Query '{sheet_name}' returned zero records.")

            if not results:
                st.warning("No records found. No file generated.")
                st.session_state["generated_results"] = None
                return

            if output_mode == "One Excel file (multiple sheets)":
                wb = openpyxl.Workbook()
                wb.remove(wb.active)

                for sheet_name, df in results:
                    ws = wb.create_sheet(title=sheet_name[:31])
                    for col_idx, col_name in enumerate(df.columns, 1):
                        ws.cell(row=1, column=col_idx, value=col_name)
                    for row_idx, row_values in enumerate(df.values, 2):
                        for col_idx, value in enumerate(row_values, 1):
                            if pd.isna(value):
                                cell_value = "NULL"
                            elif str(value).strip() == "":
                                cell_value = None
                            else:
                                cell_value = _sanitize(value)
                            ws.cell(
                                row=row_idx, column=col_idx, value=cell_value
                            )

                buf = BytesIO()
                wb.save(buf)
                file_payload = buf.getvalue()

            else:
                file_payload = [
                    (sheet_name, _build_excel_file(sheet_name, df))
                    for sheet_name, df in results
                ]

            st.session_state["generated_results"] = {
                "mode": output_mode,
                "payload": file_payload,
            }

            elapsed = time.perf_counter() - t_start
            log_event(
                feature="File Generator",
                duration_sec=elapsed,
                details={
                    "files_count": len(results),
                    "items": items_summary,
                    "total_records": sum(it["rows"] for it in items_summary),
                },
            )

            st.success(f"Processed {len(results)} queries successfully in {elapsed:.2f}s.")

        except Exception as e:
            st.error(f"Error executing queries: {e}")
            st.session_state["generated_results"] = None

    if st.session_state["generated_results"] is not None:
        saved_mode = st.session_state["generated_results"]["mode"]
        saved_payload = st.session_state["generated_results"]["payload"]

        st.divider()
        st.subheader("Generated Exports")

        if saved_mode == "One Excel file (multiple sheets)":
            st.download_button(
                "⬇ Download Combined Excel",
                data=saved_payload,
                file_name="DQ_Validation_Export.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        else:
            st.info(
                f"Generated {len(saved_payload)} distinct workbook files. Download them below:"
            )
            cols = st.columns(2)

            for i, (sheet_name, excel_bytes) in enumerate(saved_payload):
                with cols[i % 2]:
                    st.download_button(
                        label=f"⬇ Download {sheet_name}.xlsx",
                        data=excel_bytes,
                        file_name=f"{sheet_name}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"dl_btn_{sheet_name}_{i}",
                    )


if __name__ == "__main__":
    run()