import time
import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from conn_picker import render_picker
from db_utils import engine_from_profile, make_safe_table_name
from metrics_tracker import log_event


def run():
    st.title("🔍 Schema Validator")
    st.caption("Compare columns, datatypes, row counts and nulls between source and target")

    st.subheader("Connections")
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Source**")
        src_profile = render_picker("Source connection", key="schema_src")
    with col2:
        st.markdown("**Target**")
        trg_profile = render_picker("Target connection", key="schema_trg")

    with st.expander("Renamed table mappings (optional)"):
        st.caption(
            "Use if source and target tables have different names. "
            "One mapping per line: `source_schema.source_table:target_schema.target_table`"
        )
        mapping_text = st.text_area(
            "Mappings",
            height=100,
            placeholder="STG.STG_SALES:DWH.FACT_SALES\nODS.SFDC_RECORD:DWH.DIM_RECORD",
        )

    table_mapping: dict = {}
    if mapping_text:
        for line in mapping_text.strip().split("\n"):
            line = line.strip()
            if ":" in line:
                src_part, tgt_part = line.split(":", 1)
                src_part = src_part.strip().upper()
                tgt_part = tgt_part.strip().upper()
                if "." in src_part and "." in tgt_part:
                    table_mapping[src_part] = tgt_part

    st.subheader("Table fetch query")
    default_q = "SELECT TABLE_SCHEMA, TABLE_NAME FROM INFORMATION_SCHEMA.TABLES WHERE TABLE_TYPE = 'BASE TABLE'"
    use_same = st.checkbox("Use same query for source and target", value=True)
    src_query = st.text_area("Source query", value=default_q, height=100)
    trg_query = src_query if use_same else st.text_area("Target query", value=default_q, height=100)

    st.subheader("Validation Settings")
    col_check = st.checkbox("Include column & datatype validation", value=True)
    null_check = st.checkbox("Include null count validation", value=True)

    variance_threshold = st.number_input(
        "Highlighting Variance Threshold (%)",
        min_value=0.0, value=1.0, step=0.1,
        help="Rows exceeding this percentage variance will be highlighted in Red Font."
    )

    st.divider()

    if st.button("▶ Run Validation", type="primary"):
        if not src_profile or not trg_profile:
            st.error("Please select both source and target connections.")
            return

        t_start = time.perf_counter()
        try:
            src_engine = engine_from_profile(src_profile)
            trg_engine = engine_from_profile(trg_profile)

            def get_tables(engine, query):
                df = pd.read_sql(query, engine)
                df["FULL_NAME"] = df["TABLE_SCHEMA"].str.upper() + "." + df["TABLE_NAME"].str.upper()
                return set(df["FULL_NAME"].tolist())

            def table_exists(engine, schema_table):
                schema, table = schema_table.split(".", 1)
                sql = (f"SELECT COUNT(*) CNT FROM INFORMATION_SCHEMA.TABLES "
                       f"WHERE UPPER(TABLE_SCHEMA)='{schema.upper()}' AND UPPER(TABLE_NAME)='{table.upper()}'")
                return pd.read_sql(sql, engine).iloc[0]["CNT"] > 0

            def get_metadata(engine, schema, table):
                sql = (f"SELECT COLUMN_NAME, UPPER(DATA_TYPE) DATA_TYPE, "
                       f"CHARACTER_MAXIMUM_LENGTH CHAR_LENGTH, "
                       f"NUMERIC_PRECISION DATA_PRECISION, NUMERIC_SCALE DATA_SCALE "
                       f"FROM INFORMATION_SCHEMA.COLUMNS "
                       f"WHERE TABLE_SCHEMA='{schema}' AND TABLE_NAME='{table}'")
                return pd.read_sql(sql, engine)

            src_tables = get_tables(src_engine, src_query)
            trg_tables = get_tables(trg_engine, trg_query)

            matching, missing_target, missing_source = [], [], []

            for s, t in table_mapping.items():
                if table_exists(src_engine, s) and table_exists(trg_engine, t):
                    matching.append((s, t))

            already_src = {s for s, _ in matching}
            already_trg = {t for _, t in matching}
            for s in src_tables - already_src:
                if s in trg_tables:
                    matching.append((s, s))
                else:
                    missing_target.append(s)
            for t in trg_tables - already_trg:
                if t not in [x for _, x in matching]:
                    missing_source.append(t)

            if not matching:
                st.error("No matching tables found.")
                return

            progress = st.progress(0, text="Validating tables…")
            total = len(matching)

            row_data = []
            for i, (s, t) in enumerate(matching):
                s_cnt = pd.read_sql(f"SELECT COUNT(*) CNT FROM {make_safe_table_name(s)}", src_engine).iloc[0]["CNT"]
                t_cnt = pd.read_sql(f"SELECT COUNT(*) CNT FROM {make_safe_table_name(t)}", trg_engine).iloc[0]["CNT"]
                diff = abs(s_cnt - t_cnt)
                pct = (diff / max(s_cnt, t_cnt)) * 100 if max(s_cnt, t_cnt) else 0
                row_data.append([s, t, s_cnt, t_cnt, diff, round(pct, 2)])
                progress.progress((i + 1) / total * 0.3, text=f"Row count: {s}")

            df_rows = pd.DataFrame(row_data, columns=[
                "SOURCE TABLE", "TARGET TABLE", "SOURCE ROWS", "TARGET ROWS", "DIFF", "% VARIANCE"
            ])

            diff_cols = []
            missing_src_cols_list = []
            missing_tgt_cols_list = []

            if col_check:
                for i, (s, t) in enumerate(matching):
                    sch_s, tab_s = s.split(".")
                    sch_t, tab_t = t.split(".")
                    df_s = get_metadata(src_engine, sch_s, tab_s)
                    df_t = get_metadata(trg_engine, sch_t, tab_t)
                    comp = pd.merge(df_s, df_t, on="COLUMN_NAME", how="outer", suffixes=("_SRC", "_TGT"))

                    for _, r in comp.iterrows():
                        col = str(r["COLUMN_NAME"]).upper()
                        if pd.isna(r.get("DATA_TYPE_SRC")):
                            missing_src_cols_list.append([tab_s.upper(), col])
                        elif pd.isna(r.get("DATA_TYPE_TGT")):
                            missing_tgt_cols_list.append([tab_t.upper(), col])
                        else:
                            def fmt(a, b):
                                if pd.isna(a) or str(a).upper() == "NAN":
                                    a_str = "NULL"
                                elif str(a).strip() == "":
                                    a_str = "BLANK"
                                else:
                                    a_str = str(a).strip()

                                if pd.isna(b) or str(b).upper() == "NAN":
                                    b_str = "NULL"
                                elif str(b).strip() == "":
                                    b_str = "BLANK"
                                else:
                                    b_str = str(b).strip()

                                return f"{a_str} → {b_str}" if a_str != b_str else a_str

                            diff_cols.append({
                                "TABLE NAME": tab_s.upper(), "COLUMN NAME": col,
                                "DATA TYPE": fmt(r["DATA_TYPE_SRC"], r["DATA_TYPE_TGT"]),
                                "CHAR LENGTH": fmt(r.get("CHAR_LENGTH_SRC"), r.get("CHAR_LENGTH_TGT")),
                                "PRECISION": fmt(r.get("DATA_PRECISION_SRC"), r.get("DATA_PRECISION_TGT")),
                                "SCALE": fmt(r.get("DATA_SCALE_SRC"), r.get("DATA_SCALE_TGT")),
                            })
                    progress.progress(0.3 + ((i + 1) / total * 0.3), text=f"Datatype check: {s}")

            df_dtype = pd.DataFrame(diff_cols) if diff_cols else pd.DataFrame(
                columns=["TABLE NAME", "COLUMN NAME", "DATA TYPE", "CHAR LENGTH", "PRECISION", "SCALE"])
            df_missing_src_cols = pd.DataFrame(missing_src_cols_list, columns=["Table Name", "Column Name"])
            df_missing_tgt_cols = pd.DataFrame(missing_tgt_cols_list, columns=["Table Name", "Column Name"])

            null_records = []
            if null_check:
                for i, (s, t) in enumerate(matching):
                    sch_s, tab_s = s.split(".")
                    sch_t, tab_t = t.split(".")
                    cols_list = get_metadata(src_engine, sch_s, tab_s)["COLUMN_NAME"].tolist()
                    for col in cols_list:
                        s_null = pd.read_sql(
                            f"SELECT COUNT(*) CNT FROM {make_safe_table_name(s)} WHERE (CASE WHEN [{col}] IS NULL THEN 1 ELSE 0 END) = 1",
                            src_engine).iloc[0]["CNT"]
                        s_blank = pd.read_sql(
                            f"SELECT COUNT(*) CNT FROM {make_safe_table_name(s)} WHERE (CASE WHEN CAST([{col}] AS VARCHAR) = '' THEN 1 ELSE 0 END) = 1",
                            src_engine).iloc[0]["CNT"]
                        s_sum = s_null + s_blank

                        t_null = pd.read_sql(
                            f"SELECT COUNT(*) CNT FROM {make_safe_table_name(t)} WHERE (CASE WHEN [{col}] IS NULL THEN 1 ELSE 0 END) = 1",
                            trg_engine).iloc[0]["CNT"]
                        t_blank = pd.read_sql(
                            f"SELECT COUNT(*) CNT FROM {make_safe_table_name(t)} WHERE (CASE WHEN CAST([{col}] AS VARCHAR) = '' THEN 1 ELSE 0 END) = 1",
                            trg_engine).iloc[0]["CNT"]
                        t_sum = t_null + t_blank

                        diff = abs(s_sum - t_sum)
                        pct = (diff / max(s_sum, t_sum)) * 100 if max(s_sum, t_sum) else 0

                        null_records.append([
                            tab_s.upper(), col.upper(), s_blank, s_null, s_sum,
                            tab_t.upper(), col.upper(), t_blank, t_null, t_sum,
                            diff, round(pct, 2)
                        ])
                    progress.progress(0.6 + ((i + 1) / total * 0.4), text=f"Null check: {s}")

            null_headers = [
                "SOURCE TABLE_NAME", "SOURCE COLUMN_NAME", "SOURCE BLANK COUNT", "SOURCE NULL COUNT",
                "SOURCE SUM OF NULL AND BLANK",
                "TARGET TABLE_NAME", "TARGET COLUMN_NAME", "TARGET BLANK COUNT", "TARGET NULL COUNT",
                "TARGET SUM OF NULL AND BLANK",
                "DIFF", "PERCENTAGE"
            ]
            df_null = pd.DataFrame(null_records, columns=null_headers) if null_records else pd.DataFrame(
                columns=null_headers)

            progress.empty()

            elapsed = time.perf_counter() - t_start
            log_event(
                feature="Schema Validator",
                duration_sec=elapsed,
                details={
                    "tables_count": len(matching),
                    "column_check": col_check,
                    "null_check": null_check,
                    "row_count_check": True,
                },
            )

            st.subheader("Results")

            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Tables matched", len(matching))
            m2.metric("Missing in target", len(missing_target))
            m3.metric("Missing in source", len(missing_source))
            m4.metric("Row count issues", int((df_rows["% VARIANCE"] > variance_threshold).sum()))

            ui_tabs = ["Table Status"]
            if col_check:
                ui_tabs.extend(["Column Status", "Datatype check"])
            ui_tabs.append("Row counts")
            if null_check:
                ui_tabs.append("Null check")

            rendered_tabs = st.tabs(ui_tabs)
            curr_tab = 0

            with rendered_tabs[curr_tab]:
                st.dataframe(df_rows[["SOURCE TABLE", "TARGET TABLE"]], use_container_width=True, hide_index=True)
            curr_tab += 1

            if col_check:
                with rendered_tabs[curr_tab]:
                    st.markdown("**Missing Columns in Source:**")
                    st.dataframe(df_missing_src_cols, use_container_width=True, hide_index=True)
                    st.markdown("**Missing Columns in Target:**")
                    st.dataframe(df_missing_tgt_cols, use_container_width=True, hide_index=True)
                curr_tab += 1
                with rendered_tabs[curr_tab]:
                    st.dataframe(df_dtype, use_container_width=True, hide_index=True)
                curr_tab += 1

            with rendered_tabs[curr_tab]:
                st.dataframe(df_rows, use_container_width=True, hide_index=True)
            curr_tab += 1

            if null_check:
                with rendered_tabs[curr_tab]:
                    st.dataframe(df_null, use_container_width=True, hide_index=True)

            buffer = BytesIO()
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thick_blue_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
            yellow_diff_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            red_alert_font = Font(name="Calibri", size=11, bold=True, color="FF0000")

            center_align = Alignment(horizontal="center", vertical="center")
            left_align = Alignment(horizontal="left", vertical="center")

            with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
                wb = writer.book

                ws1 = wb.create_sheet(title="Table Check")
                match_series = [f"{s} -> {t}" for s, t in matching]
                max_t_len = max(len(match_series), len(missing_source), len(missing_target), 1)

                match_series += [""] * (max_t_len - len(match_series))
                missing_source_extended = missing_source + [""] * (max_t_len - len(missing_source))
                missing_target_extended = missing_target + [""] * (max_t_len - len(missing_target))

                df_table_check = pd.DataFrame({
                    "Matching Tables": match_series,
                    "Missing Tables in Source": missing_source_extended,
                    "Missing Tables in Target": missing_target_extended
                })

                for col_idx, col_name in enumerate(df_table_check.columns, 1):
                    cell = ws1.cell(row=1, column=col_idx, value=col_name)
                    cell.font = header_font
                    cell.fill = thick_blue_fill
                    cell.alignment = center_align

                for row_idx, row_values in enumerate(df_table_check.itertuples(index=False), 2):
                    for col_idx, val in enumerate(row_values, 1):
                        ws1.cell(row=row_idx, column=col_idx, value=val).alignment = left_align

                if col_check:
                    ws2 = wb.create_sheet(title="Column Check")
                    c1 = ws2.cell(row=1, column=1, value="COLUMNS MISSING IN SOURCE")
                    c1.font = Font(name="Calibri", bold=True, size=11)
                    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)

                    for c_idx, text in enumerate(["Table Name", "Column Name"], 1):
                        sub_cell = ws2.cell(row=2, column=c_idx, value=text)
                        sub_cell.font = header_font
                        sub_cell.fill = thick_blue_fill

                    current_r = 3
                    for row in df_missing_src_cols.itertuples(index=False):
                        ws2.cell(row=current_r, column=1, value=row[0])
                        ws2.cell(row=current_r, column=2, value=row[1])
                        current_r += 1

                    current_r += 2
                    c2 = ws2.cell(row=current_r, column=1, value="COLUMNS MISSING IN TARGET")
                    c2.font = Font(name="Calibri", bold=True, size=11)
                    ws2.merge_cells(start_row=current_r, start_column=1, end_row=current_r, end_column=2)

                    current_r += 1
                    for c_idx, text in enumerate(["Table Name", "Column Name"], 1):
                        sub_cell = ws2.cell(row=current_r, column=c_idx, value=text)
                        sub_cell.font = header_font
                        sub_cell.fill = thick_blue_fill

                    current_r += 1
                    for row in df_missing_tgt_cols.itertuples(index=False):
                        ws2.cell(row=current_r, column=1, value=row[0])
                        ws2.cell(row=current_r, column=2, value=row[1])
                        current_r += 1

                if col_check:
                    ws3 = wb.create_sheet(title="DataType Check")
                    for col_idx, col_name in enumerate(df_dtype.columns, 1):
                        cell = ws3.cell(row=1, column=col_idx, value=col_name)
                        cell.font = header_font
                        cell.fill = thick_blue_fill
                        cell.alignment = center_align

                    for row_idx, row_values in enumerate(df_dtype.itertuples(index=False), 2):
                        for col_idx, val in enumerate(row_values, 1):
                            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
                            if "→" in str(val or ""):
                                cell.fill = yellow_diff_fill

                ws4 = wb.create_sheet(title="Row Count Check")
                for col_idx, col_name in enumerate(df_rows.columns, 1):
                    cell = ws4.cell(row=1, column=col_idx, value=col_name)
                    cell.font = header_font
                    cell.fill = thick_blue_fill
                    cell.alignment = center_align

                for row_idx, row_values in enumerate(df_rows.itertuples(index=False), 2):
                    for col_idx, val in enumerate(row_values, 1):
                        cell = ws4.cell(row=row_idx, column=col_idx, value=val)
                        if col_idx == 6 and isinstance(val, (int, float)) and val > variance_threshold:
                            cell.font = red_alert_font

                if null_check:
                    ws5 = wb.create_sheet(title="Null Count Check")
                    for col_idx, col_name in enumerate(df_null.columns, 1):
                        cell = ws5.cell(row=1, column=col_idx, value=col_name)
                        cell.font = header_font
                        cell.fill = thick_blue_fill
                        cell.alignment = center_align

                    for row_idx, row_values in enumerate(df_null.itertuples(index=False), 2):
                        for col_idx, val in enumerate(row_values, 1):
                            cell = ws5.cell(row=row_idx, column=col_idx, value=val)
                            if col_idx == 12 and isinstance(val, (int, float)) and val > variance_threshold:
                                cell.font = red_alert_font

                for sheet in wb.worksheets:
                    for col in sheet.columns:
                        vals = [len(str(cell.value or '')) for cell in col]
                        max_len = max(vals) if vals else 10
                        col_letter = get_column_letter(col[0].column)
                        sheet.column_dimensions[col_letter].width = max(max_len + 3, 18)

            if "Sheet" in wb.sheetnames:
                wb.remove(wb["Sheet"])

            buffer.seek(0)
            st.download_button(
                "⬇ Download Validation Report",
                buffer,
                "SCHEMA_VALIDATION_REPORT.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        except Exception as e:
            st.error(f"Error compiling validation report metrics: {e}")


if __name__ == "__main__":
    run()