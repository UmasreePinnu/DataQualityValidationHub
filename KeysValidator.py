import time
import streamlit as st
import pandas as pd
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font
from conn_picker import render_picker
from db_utils import engine_from_profile
from metrics_tracker import log_event


def _get_key_cols(engine, schema, table):
    """Extract columns ending in _KEY (case-insensitive) from target table metadata."""
    try:
        clean_schema = schema.strip().upper()
        clean_table = table.strip().upper()

        q = f"""
            SELECT COLUMN_NAME 
            FROM INFORMATION_SCHEMA.COLUMNS 
            WHERE UPPER(TABLE_SCHEMA) = '{clean_schema}' 
              AND UPPER(TABLE_NAME) = '{clean_table}' 
              AND UPPER(COLUMN_NAME) LIKE '%_KEY'
        """
        df = pd.read_sql(q, engine)
        if df.empty:
            q_fallback = f"""
                SELECT COLUMN_NAME 
                FROM INFORMATION_SCHEMA.COLUMNS 
                WHERE UPPER(TABLE_NAME) = '{clean_table}' 
                  AND UPPER(COLUMN_NAME) LIKE '%_KEY'
            """
            df = pd.read_sql(q_fallback, engine)

        return [c.upper() for c in df["COLUMN_NAME"].tolist()]
    except Exception as e:
        st.error(f"Metadata fetch error: {e}")
        return []


def run():
    st.title("🔑 Keys Validator")
    st.caption("Validate referential integrity, duplicate keys and null counts")

    scenario = st.radio(
        "Validation scenario",
        ["Fact → Dim integrity", "Dim → Fact integrity", "Duplicate check", "Null count"],
        horizontal=True,
        label_visibility="collapsed",
    )

    st.divider()

    # ══════════════════════════════════════════════════════════════════════════
    # SCENARIOS 1 & 2: KEY INTEGRITY
    # ══════════════════════════════════════════════════════════════════════════
    if scenario in ["Fact → Dim integrity", "Dim → Fact integrity"]:
        fact_first = scenario == "Fact → Dim integrity"
        label_main = "Fact table" if fact_first else "Dimension table"
        label_map = "Dimension table" if fact_first else "Fact table"

        st.subheader("Connection & environment")
        c1, c2 = st.columns(2)
        with c1:
            env_profile = render_picker("Environment", key="keys_env")
        with c2:
            st.markdown(f"**{label_main}**")
            main_table = st.text_input(
                label_main,
                placeholder="schema.table_name",
                label_visibility="collapsed",
            )

        if env_profile and main_table and "." in main_table:
            engine = engine_from_profile(env_profile)
            schema, table = main_table.split(".", 1)
            keys = _get_key_cols(engine, schema, table)

            if keys:
                st.subheader(f"Map _KEY columns → {label_map}s")
                mappings = {}
                cols = st.columns(2)
                for i, key in enumerate(keys):
                    suggested = (
                        f"DWH.DIM_{key.replace('_KEY', '')}"
                        if fact_first
                        else f"DWH.FACT_{key.replace('_KEY', '')}"
                    )
                    mappings[key] = cols[i % 2].text_input(
                        f"{label_map} for {key}",
                        value=suggested,
                        key=f"key_map_{key}",
                    )

                if st.button("▶ Run Validation", type="primary"):
                    t_start = time.perf_counter()
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "SUMMARY"
                    ws.append(["TABLE NAME", "KEY / WID NAME", "STATUS", "QUERY USED"])

                    progress = st.progress(0)
                    for ki, (key, dim_table) in enumerate(mappings.items()):
                        if not dim_table:
                            continue

                        if fact_first:
                            sql = (
                                f"SELECT * FROM {main_table} f "
                                f"WHERE NOT EXISTS (SELECT 1 FROM {dim_table} d WHERE f.{key}=d.{key}) "
                                f"AND f.{key} IS NOT NULL"
                            )
                            status_lbl = "Keys present in Fact not in Dim Table"
                        else:
                            sql = (
                                f"SELECT d.{key} FROM {dim_table} d "
                                f"WHERE NOT EXISTS (SELECT 1 FROM {main_table} f WHERE d.{key}=f.{key})"
                            )
                            status_lbl = "Keys present in Dim not in Fact Table"

                        try:
                            df_err = pd.read_sql(sql, engine)
                            status = status_lbl if not df_err.empty else "no issues"
                            ws.append([main_table.upper(), key.upper(), status, sql])

                            ws_det = wb.create_sheet(title=f"Missing_{key}"[:30])
                            if df_err.empty:
                                ws_det.append(["no issues"])
                            else:
                                ws_det.append([c.upper() for c in df_err.columns])
                                for r in df_err.values.tolist():
                                    ws_det.append(r)
                        except Exception as e:
                            st.error(f"Error on {key}: {e}")

                        progress.progress((ki + 1) / len(mappings))

                    progress.empty()
                    buf = BytesIO()
                    wb.save(buf)

                    elapsed = time.perf_counter() - t_start
                    log_event(
                        feature="Keys Validator",
                        duration_sec=elapsed,
                        details={
                            "scenario": scenario,
                            "validations_count": len(mappings),
                        },
                    )

                    st.success("Validation complete.")
                    st.download_button(
                        "⬇ Download Report",
                        buf.getvalue(),
                        f"Key_Integrity_{datetime.now().strftime('%Y%m%d')}.xlsx",
                    )
            else:
                st.info("No _KEY columns detected. Enter a valid table with _KEY columns.")

    # ══════════════════════════════════════════════════════════════════════════
    # SCENARIO 3: DUPLICATE CHECK
    # ══════════════════════════════════════════════════════════════════════════
    elif scenario == "Duplicate check":
        st.subheader("Connection & table")
        c1, c2 = st.columns(2)
        with c1:
            env_profile = render_picker("Environment", key="dup_env")
        with c2:
            target_tbl = st.text_input("Table name", placeholder="schema.table_name")

        pk_cols = st.text_input(
            "Primary key columns (comma separated)",
            placeholder="COL1, COL2, COL3",
        )

        if st.button("▶ Check Duplicates", type="primary"):
            if not env_profile:
                st.error("Please select a connection.")
                return
            if not target_tbl or not pk_cols:
                st.error("Table name and key columns are required.")
                return

            t_start = time.perf_counter()
            engine = engine_from_profile(env_profile)
            sql = (
                f"SELECT {pk_cols}, COUNT(*) AS DUP_COUNT "
                f"FROM {target_tbl} "
                f"GROUP BY {pk_cols} HAVING COUNT(*) > 1"
            )
            try:
                df_dups = pd.read_sql(sql, engine)
                elapsed = time.perf_counter() - t_start

                log_event(
                    feature="Keys Validator",
                    duration_sec=elapsed,
                    details={
                        "scenario": "Duplicate check",
                        "validations_count": 1,
                    },
                )

                if df_dups.empty:
                    st.success("✅ No duplicates found.")
                else:
                    st.warning(f"⚠️ Found **{len(df_dups)}** duplicate combination(s).")
                    st.dataframe(df_dups, use_container_width=True, hide_index=True)

                    wb = Workbook()
                    ws = wb.active
                    ws.title = "DUPLICATES"
                    ws.append([c.upper() for c in df_dups.columns])
                    for r in df_dups.values.tolist():
                        ws.append(r)
                    buf = BytesIO()
                    wb.save(buf)
                    st.download_button(
                        "⬇ Download Duplicate Report",
                        buf.getvalue(),
                        "Duplicate_Report.xlsx",
                    )
            except Exception as e:
                st.error(f"SQL error: {e}")

    # ══════════════════════════════════════════════════════════════════════════
    # SCENARIO 4: NULL COUNT
    # ══════════════════════════════════════════════════════════════════════════
    elif scenario == "Null count":
        st.subheader("Connections")
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Source**")
            src_profile = render_picker("Source", key="null_src")
        with c2:
            st.markdown("**Target**")
            trg_profile = render_picker("Target", key="null_trg")

        st.divider()
        st.subheader("Table pairs")

        if "null_pairs" not in st.session_state:
            st.session_state["null_pairs"] = [{"src": "", "trg": ""}]

        for i, pair in enumerate(st.session_state["null_pairs"]):
            ca, cb, cc = st.columns([4, 4, 0.5])
            pair["src"] = ca.text_input(
                "Source table",
                value=pair["src"],
                key=f"np_src_{i}",
                placeholder="STG.FACT_SALES",
            )
            pair["trg"] = cb.text_input(
                "Target table",
                value=pair["trg"],
                key=f"np_trg_{i}",
                placeholder="DWH.FACT_SALES",
            )
            if cc.button("×", key=f"np_del_{i}") and len(st.session_state["null_pairs"]) > 1:
                st.session_state["null_pairs"].pop(i)
                st.rerun()

        if st.button("＋ Add table pair"):
            st.session_state["null_pairs"].append({"src": "", "trg": ""})
            st.rerun()

        if st.button("▶ Run Null Comparison", type="primary"):
            if not src_profile or not trg_profile:
                st.error("Please select both source and target connections.")
                return

            t_start = time.perf_counter()
            s_engine = engine_from_profile(src_profile)
            t_engine = engine_from_profile(trg_profile)

            wb = Workbook()
            ws = wb.active
            ws.title = "NULL_COUNT_CHECK"
            ws.append([
                "TABLE NAME",
                "COLUMN NAME",
                "SOURCE NULL COUNT",
                "TARGET NULL COUNT",
                "DIFF",
                "% VARIANCE",
            ])

            records = []
            pairs = [p for p in st.session_state["null_pairs"] if p["src"] and p["trg"]]

            if not pairs:
                st.error("Please specify at least one valid table pair.")
                return

            progress = st.progress(0)

            for pi, pair in enumerate(pairs):
                trg_schema, trg_table = (
                    pair["trg"].split(".", 1)
                    if "." in pair["trg"]
                    else ("", pair["trg"])
                )
                keys = _get_key_cols(t_engine, trg_schema, trg_table)

                for k in keys:
                    try:
                        sc = pd.read_sql(
                            f"SELECT COUNT(*) AS C FROM {pair['src']} WHERE {k} IS NULL",
                            s_engine,
                        ).iloc[0, 0]
                        tc = pd.read_sql(
                            f"SELECT COUNT(*) AS C FROM {pair['trg']} WHERE {k} IS NULL",
                            t_engine,
                        ).iloc[0, 0]
                        diff = abs(sc - tc)
                        var = (diff / sc * 100) if sc > 0 else (0 if diff == 0 else 100)
                        var = round(var, 2)

                        records.append({
                            "TABLE": pair["trg"].upper(),
                            "COLUMN": k.upper(),
                            "SRC NULLS": sc,
                            "TGT NULLS": tc,
                            "DIFF": diff,
                            "% VARIANCE": var,
                        })
                        ws.append([pair["trg"].upper(), k.upper(), sc, tc, diff, var])

                        if var > 1:
                            for cell in ws[ws.max_row]:
                                cell.font = Font(color="FF0000")
                    except Exception:
                        continue

                progress.progress((pi + 1) / len(pairs))

            progress.empty()

            elapsed = time.perf_counter() - t_start
            log_event(
                feature="Keys Validator",
                duration_sec=elapsed,
                details={
                    "scenario": "Null count",
                    "validations_count": len(pairs),
                },
            )

            if records:
                df_res = pd.DataFrame(records)

                def highlight_variance(row):
                    return [
                        "color: red; font-weight: bold" if row["% VARIANCE"] > 1 else ""
                        for _ in row
                    ]

                st.dataframe(
                    df_res.style.apply(highlight_variance, axis=1),
                    use_container_width=True,
                    hide_index=True,
                )

                buf = BytesIO()
                wb.save(buf)
                st.success("Null comparison complete.")
                st.download_button(
                    "⬇ Download Null Report",
                    buf.getvalue(),
                    "Null_Count_Report.xlsx",
                )
            else:
                st.info("No _KEY columns found in the target tables or queries returned no results.")


if __name__ == "__main__":
    run()