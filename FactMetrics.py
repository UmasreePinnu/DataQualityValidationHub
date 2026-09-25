import time
import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from conn_picker import render_picker
from db_utils import engine_from_profile
from metrics_tracker import log_event

AGGREGATE_OPTIONS = ["SUM", "AVG", "MIN", "MAX", "COUNT"]

_FILL_SRC   = PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid")
_FILL_TGT   = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
_FILL_DIFF  = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
_FILL_PERC  = PatternFill(start_color="87CEFA", end_color="87CEFA", fill_type="solid")
_FILL_AGG   = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
_FILL_ALERT = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

_FONT_WHITE     = Font(bold=True, color="FFFFFF")
_FONT_RED_BOLD  = Font(bold=True, color="FF0000")
_FONT_BOLD      = Font(bold=True)


def _numeric_cols(engine, schema, table):
    q = (
        f"SELECT COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS "
        f"WHERE TABLE_SCHEMA='{schema}' AND UPPER(TABLE_NAME)='{table.upper()}' "
        f"AND DATA_TYPE IN ('int','bigint','decimal','float','numeric','real','smallint') "
        f"AND COLUMN_NAME NOT LIKE '%[_]KEY' AND COLUMN_NAME NOT LIKE '%[_]WID'"
    )
    return [c.upper() for c in pd.read_sql(q, engine)["COLUMN_NAME"].tolist()]


def _pct(src, tgt):
    src = float(src or 0)
    tgt = float(tgt or 0)
    if src == 0 and tgt == 0:
        return 0.0
    if src == 0:
        return 100.0
    return round(abs((src - tgt) / src) * 100, 4)


def _compare_one_mapping(src_engine, trg_engine, src_t, tgt_t, dims, aggs):
    sch_s, tab_s = src_t.split(".", 1)
    sch_t, tab_t = tgt_t.split(".", 1)

    src_cols    = _numeric_cols(src_engine, sch_s, tab_s)
    tgt_cols    = _numeric_cols(trg_engine, sch_t, tab_t)
    common_cols = [c for c in src_cols if c in tgt_cols]

    if not common_cols:
        raise ValueError("No common numeric metric columns found between source and target.")

    dim_select = (", ".join(dims) + ", ") if dims else ""
    grp        = (f" GROUP BY {', '.join(dims)}") if dims else ""

    agg_dfs = {}

    for agg in aggs:
        src_sel = ", ".join(f"{agg}({c}) AS {agg}_{c}_SRC" for c in common_cols)
        tgt_sel = ", ".join(f"{agg}({c}) AS {agg}_{c}_TGT" for c in common_cols)

        df_src = pd.read_sql(f"SELECT {dim_select}{src_sel} FROM {src_t}{grp}", src_engine)
        df_tgt = pd.read_sql(f"SELECT {dim_select}{tgt_sel} FROM {tgt_t}{grp}", trg_engine)

        df_src.columns = [c.upper() for c in df_src.columns]
        df_tgt.columns = [c.upper() for c in df_tgt.columns]

        use_dummy  = not dims
        merge_keys = [d.upper() for d in dims] if dims else ["_MERGE_KEY"]
        if use_dummy:
            df_src["_MERGE_KEY"] = 1
            df_tgt["_MERGE_KEY"] = 1

        df_m = pd.merge(df_src, df_tgt, on=merge_keys, how="outer", suffixes=("", ""))

        for col in df_m.columns:
            if col.endswith("_SRC") or col.endswith("_TGT"):
                df_m[col] = pd.to_numeric(df_m[col], errors="coerce").fillna(0)

        final_cols = [k for k in merge_keys if k != "_MERGE_KEY"]
        for metric in common_cols:
            sc = f"{agg}_{metric}_SRC"
            tc = f"{agg}_{metric}_TGT"
            dc = f"DIFF_{agg}_{metric}"
            pc = f"PERC_{agg}_{metric}"
            if sc in df_m.columns and tc in df_m.columns:
                df_m[dc] = df_m[sc] - df_m[tc]
                df_m[pc] = df_m.apply(lambda r, s=sc, t=tc: _pct(r[s], r[t]), axis=1)
                final_cols += [sc, tc, dc, pc]

        if use_dummy and "_MERGE_KEY" in df_m.columns:
            df_m.drop(columns=["_MERGE_KEY"], inplace=True)

        agg_dfs[agg] = df_m[[c for c in final_cols if c in df_m.columns]]

    return agg_dfs, tab_s, tab_t


def _autofit(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)


def _build_details_df(results, threshold):
    rows = []
    for pair_label, agg_dfs in results.items():
        for agg, df in agg_dfs.items():
            perc_cols = [c for c in df.columns if c.startswith(f"PERC_{agg}_")]
            for _, row in df.iterrows():
                for pc in perc_cols:
                    pct_val = row[pc]
                    if not isinstance(pct_val, (int, float)):
                        continue
                    if pct_val <= threshold:
                        continue
                    metric = pc[len(f"PERC_{agg}_"):]
                    sc = f"{agg}_{metric}_SRC"
                    tc = f"{agg}_{metric}_TGT"
                    dc = f"DIFF_{agg}_{metric}"
                    rows.append({
                        "TABLE PAIR":     pair_label,
                        "AGGREGATE":      agg,
                        "METRIC":         metric,
                        "SOURCE VALUE":   row.get(sc, ""),
                        "TARGET VALUE":   row.get(tc, ""),
                        "DIFF":           row.get(dc, ""),
                        "% VARIANCE":     round(pct_val, 4),
                    })
    return pd.DataFrame(rows)


def run():
    st.title("📊 Fact Metrics Validator")
    st.caption("Compare numeric metric aggregates between source and target tables")

    st.subheader("Connections")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Source**")
        src_profile = render_picker("Source connection", key="fact_src")
    with c2:
        st.markdown("**Target**")
        trg_profile = render_picker("Target connection", key="fact_trg")

    st.divider()

    st.subheader("Aggregate functions")
    st.caption("Select one or more — results for each are stacked in the output.")

    agg_cols      = st.columns(len(AGGREGATE_OPTIONS))
    selected_aggs = []
    for col, agg in zip(agg_cols, AGGREGATE_OPTIONS):
        default = agg == "SUM"
        if col.checkbox(agg, value=default, key=f"agg_{agg}"):
            selected_aggs.append(agg)

    if not selected_aggs:
        st.warning("Select at least one aggregate function.", icon="⚠️")

    st.divider()

    st.subheader("Table mappings")
    st.caption("Optionally add dimension columns (pipe-separated) to group by.")

    if "fact_mappings" not in st.session_state:
        st.session_state["fact_mappings"] = [{"src": "", "tgt": "", "dims": ""}]

    for i, row in enumerate(st.session_state["fact_mappings"]):
        ca, cb, cc, cd = st.columns([3, 3, 3, 0.5])
        row["src"]  = ca.text_input("Source table",        value=row["src"],
                                    key=f"fm_src_{i}", placeholder="ODS.FACT_SALES")
        row["tgt"]  = cb.text_input("Target table",        value=row["tgt"],
                                    key=f"fm_tgt_{i}", placeholder="DWH.FACT_SALES")
        row["dims"] = cc.text_input("Dim columns (| sep)",  value=row["dims"],
                                    key=f"fm_dim_{i}", placeholder="FISCAL_YEAR|REGION")
        if cd.button("×", key=f"fm_del_{i}") and len(st.session_state["fact_mappings"]) > 1:
            st.session_state["fact_mappings"].pop(i)
            st.rerun()

    if st.button("＋ Add mapping"):
        st.session_state["fact_mappings"].append({"src": "", "tgt": "", "dims": ""})
        st.rerun()

    threshold = st.number_input("Variance threshold (%)", min_value=0.0, value=1.0, step=0.1)

    st.divider()

    if st.button("▶ Run Comparison", type="primary"):
        if not src_profile or not trg_profile:
            st.error("Please select both source and target connections.")
            return
        if not selected_aggs:
            st.error("Please select at least one aggregate function.")
            return

        mappings = [r for r in st.session_state["fact_mappings"] if r["src"] and r["tgt"]]
        if not mappings:
            st.error("Please enter at least one table mapping.")
            return

        t_start = time.perf_counter()
        src_engine = engine_from_profile(src_profile)
        trg_engine = engine_from_profile(trg_profile)

        all_results = {}
        progress    = st.progress(0)

        for idx, row in enumerate(mappings):
            src_t = row["src"].strip().upper()
            tgt_t = row["tgt"].strip().upper()
            dims  = [d.strip().upper() for d in row["dims"].split("|") if d.strip()] \
                    if row["dims"].strip() else []
            try:
                agg_dfs, tab_s, tab_t = _compare_one_mapping(
                    src_engine, trg_engine, src_t, tgt_t, dims, selected_aggs
                )
                key = f"{tab_s}_vs_{tab_t}"
                if key in all_results:
                    key = f"{key}_{idx}"
                all_results[key] = agg_dfs

            except Exception as e:
                st.error(f"Error on mapping {src_t} → {tgt_t}: {e}")

            progress.progress((idx + 1) / len(mappings))

        progress.empty()
        st.session_state["fact_results"]   = all_results
        st.session_state["fact_threshold"] = threshold
        st.session_state["fact_aggs"]      = selected_aggs

        elapsed = time.perf_counter() - t_start
        log_event(
            feature="Fact Metrics Validator",
            duration_sec=elapsed,
            details={
                "fact_tables_count": len(mappings),
                "table_pairs": [f"{r['src']} → {r['tgt']}" for r in mappings],
                "aggs": selected_aggs,
            },
        )

    if st.session_state.get("fact_results"):
        st.divider()
        st.subheader("Results")

        threshold = st.session_state["fact_threshold"]
        aggs      = st.session_state.get("fact_aggs", ["SUM"])
        results   = st.session_state["fact_results"]

        def highlight_row(row, thresh=threshold):
            styles = []
            for col in row.index:
                if col.startswith("PERC_") and isinstance(row[col], (int, float)) and row[col] > thresh:
                    styles.append("color:red;font-weight:bold")
                else:
                    styles.append("")
            return styles

        pair_tabs = st.tabs(list(results.keys()))
        for pair_tab, (pair_label, agg_dfs) in zip(pair_tabs, results.items()):
            with pair_tab:
                if len(aggs) == 1:
                    df = agg_dfs[aggs[0]]
                    st.dataframe(df.style.apply(highlight_row, axis=1),
                                 use_container_width=True, hide_index=True)
                else:
                    inner_tabs = st.tabs([a for a in aggs if a in agg_dfs])
                    for inner_tab, agg in zip(inner_tabs, [a for a in aggs if a in agg_dfs]):
                        with inner_tab:
                            df = agg_dfs[agg]
                            st.dataframe(df.style.apply(highlight_row, axis=1),
                                         use_container_width=True, hide_index=True)

        buffer = BytesIO()
        with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
            writer.book

            for pair_label, agg_dfs in results.items():
                sheet_name = pair_label[:31]
                writer.book.create_sheet(title=sheet_name)
                ws = writer.book[sheet_name]
                ws._current_row = 0

                first_block = True
                for agg in aggs:
                    if agg not in agg_dfs:
                        continue
                    df = agg_dfs[agg]

                    if not first_block:
                        ws.append([""])
                    first_block = False

                    label_row = ws.max_row + 1
                    for ci in range(1, df.shape[1] + 1):
                        c = ws.cell(row=label_row, column=ci,
                                    value=f"── {agg} ──" if ci == 1 else "")
                        c.font      = _FONT_WHITE
                        c.fill      = _FILL_AGG
                        c.alignment = Alignment(horizontal="left")

                    hdr_row = label_row + 1
                    for ci, col_name in enumerate(df.columns, 1):
                        cell = ws.cell(row=hdr_row, column=ci, value=col_name)
                        cell.font = _FONT_BOLD
                        cn = str(col_name)
                        if cn.endswith("_SRC"):    cell.fill = _FILL_SRC
                        elif cn.endswith("_TGT"):  cell.fill = _FILL_TGT
                        elif cn.startswith("DIFF_"): cell.fill = _FILL_DIFF
                        elif cn.startswith("PERC_"): cell.fill = _FILL_PERC

                    perc_ci = {ci for ci, cn in enumerate(df.columns, 1)
                               if cn.startswith("PERC_")}
                    for r_vals in df.itertuples(index=False):
                        data_row = ws.max_row + 1
                        for ci, val in enumerate(r_vals, 1):
                            cell = ws.cell(row=data_row, column=ci, value=val)
                            if ci in perc_ci and isinstance(val, (int, float)) and val > threshold:
                                cell.font = _FONT_RED_BOLD

                _autofit(ws)

            df_details = _build_details_df(results, threshold)

            writer.book.create_sheet(title="DETAILS")
            ws_d = writer.book["DETAILS"]

            detail_headers = ["TABLE PAIR", "AGGREGATE", "METRIC",
                               "SOURCE VALUE", "TARGET VALUE", "DIFF", "% VARIANCE"]

            for ci, h in enumerate(detail_headers, 1):
                cell = ws_d.cell(row=1, column=ci, value=h)
                cell.font      = _FONT_WHITE
                cell.fill      = _FILL_AGG
                cell.alignment = Alignment(horizontal="center")

            if df_details.empty:
                ws_d.cell(row=2, column=1,
                          value="✅ No variances exceed the threshold.")
            else:
                for ri, row in enumerate(df_details.itertuples(index=False), 2):
                    for ci, val in enumerate(row, 1):
                        cell = ws_d.cell(row=ri, column=ci, value=val)
                        if ci == 7 and isinstance(val, (int, float)) and val > threshold:
                            cell.font = _FONT_RED_BOLD

            _autofit(ws_d)
            wb = writer.book
            wb.move_sheet("DETAILS", offset=-len(wb.sheetnames) + 1)

        buffer.seek(0)
        st.download_button(
            "⬇ Download Comparison Report",
            buffer,
            "FACT_METRICS_REPORT.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


if __name__ == "__main__":
    run()