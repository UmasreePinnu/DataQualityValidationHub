import streamlit as st
import pandas as pd
import openpyxl
from io import BytesIO
from openpyxl.styles import PatternFill, Font
import time
from metrics_tracker import log_event


def format_val(v):
    """Clean and stringify values for comparison, handling floating decimals and text literals safely."""
    if pd.isna(v) or v is None:
        return ""
    val_str = str(v).strip()

    # Handle literal string missing notations safely
    if val_str.upper() in ["NA", "N/A", "NULL", "NAN"]:
        return val_str

    # Normalize float representations (e.g., '761.0' -> '761') to prevent fake numeric mismatches
    if val_str.endswith(".0"):
        base_part = val_str[:-2]
        if base_part.replace("-", "", 1).isdigit():
            return base_part

    return val_str


def get_row_key(row, cols):
    """Create a unique string key based on selected columns."""
    vals = [format_val(row[c]) for c in cols if c in row]
    return " | ".join(vals)


def is_cell_highlighted(cell):
    """Helper to verify if an openpyxl cell has a valid background fill color."""
    if cell and cell.fill and cell.fill.fill_type and cell.fill.fill_type != "none":
        if cell.fill.start_color:
            color_str = str(cell.fill.start_color.index).strip()
            if (
                color_str not in ("00000000", "64", "0", "SystemRow")
                and color_str != ""
            ):
                return True

            rgb_str = str(cell.fill.start_color.rgb).strip()
            if rgb_str not in (
                "00000000",
                "000000",
                "FFFFFFFF",
                "FFFFFF",
                "",
                "None",
            ):
                return True
    return False


def is_cell_red(cell):
    """Helper to dynamically catch all variants of red, pink, coral, and soft red theme shades in Excel."""
    if cell and cell.fill and cell.fill.fill_type and cell.fill.fill_type != "none":
        if cell.fill.start_color:
            idx_str = str(cell.fill.start_color.index).upper().strip()
            if any(
                token in idx_str
                for token in ["RED", "PINK", "CORAL", "C7CE", "FF4D4D"]
            ):
                return True

            rgb_str = str(cell.fill.start_color.rgb).upper().strip()
            if len(rgb_str) in [6, 8] and all(
                c in "0123456789ABCDEF" for c in rgb_str
            ):
                hex_val = rgb_str[-6:]
                r = int(hex_val[0:2], 16)
                g = int(hex_val[2:4], 16)
                b = int(hex_val[4:6], 16)

                if r > 160 and g < 225 and b < 225:
                    if (r - g > 30) or (r - b > 30):
                        return True

            if idx_str in [
                "FFFF0000",
                "FF0000",
                "FFC7CE",
                "9C0006",
                "10",
                "FF6666",
                "FF8080",
            ]:
                return True
    return False


def run():
    st.title("🔀 Mid File Validator")
    st.caption(
        "Compare source and reference files with row-level highlight detection and composite keys."
    )

    # 1. Feature Selection
    with st.container(border=True):
        st.subheader("1. Select Validation Mode")
        v_mode = st.radio(
            "Mode",
            ["Highlighted Values Validation", "Selected Columns Validation"],
            horizontal=True,
            label_visibility="collapsed",
        )

    # 2. Data Sources
    col1, col2 = st.columns(2)
    with col1:
        st.subheader("Source (File 1)")
        f1 = st.file_uploader(
            "Upload Source", type=["xlsx"], key="f1_uploader"
        )
        s1 = None
        if f1:
            wb1_temp = openpyxl.load_workbook(f1, read_only=True)
            s1 = st.selectbox(
                "Select Source Sheet", wb1_temp.sheetnames, key="s1_select"
            )

    with col2:
        st.subheader("Reference (File 2)")
        f2 = st.file_uploader(
            "Upload Reference", type=["xlsx"], key="f2_uploader"
        )
        s2 = None
        if f2:
            wb2_temp = openpyxl.load_workbook(f2, read_only=True)
            s2 = st.selectbox(
                "Select Reference Sheet", wb2_temp.sheetnames, key="s2_select"
            )

    if not f1 or not f2:
        st.info("Please upload both files to proceed.")
        return

    df1 = pd.read_excel(f1, sheet_name=s1, keep_default_na=False)
    df2 = pd.read_excel(f2, sheet_name=s2, keep_default_na=False)

    headers1 = [str(c) for c in df1.columns]
    headers2 = [str(c) for c in df2.columns]

    # 3. Key Configuration
    st.divider()
    st.subheader("3. Key Configuration")
    k_col1, k_col2 = st.columns(2)
    with k_col1:
        key_f1 = st.multiselect("Source Key Columns", headers1, key="kf1")
    with k_col2:
        key_f2 = st.multiselect("Reference Key Columns", headers2, key="kf2")

    # 4. Mode-specific logic
    validation_map = []
    highlighted_map = []

    wb1 = openpyxl.load_workbook(f1, data_only=True)
    ws1 = wb1[s1]

    if v_mode == "Selected Columns Validation":
        st.divider()
        st.subheader("4. Manual Column Selection")

        if "mid_v_map" not in st.session_state:
            st.session_state.mid_v_map = []

        m_col1, m_col2, m_btn = st.columns([4, 4, 1])
        with m_col1:
            map_f1 = st.selectbox("Source Col", headers1, key="mf1")
        with m_col2:
            map_f2 = st.selectbox("Reference Col", headers2, key="mf2")
        with m_btn:
            st.markdown(
                "<div style='padding-top:28px'></div>", unsafe_allow_html=True
            )
            if st.button("Add", key="add_map_btn"):
                st.session_state.mid_v_map.append({
                    "f1": map_f1,
                    "f2": map_f2,
                    "label": map_f1 if map_f1 == map_f2 else f"{map_f1} vs {map_f2}",
                })

        if st.session_state.mid_v_map:
            st.write("Active Mappings:")
            for i, mapping in enumerate(st.session_state.mid_v_map):
                if st.button(f"× {mapping['label']}", key=f"mid_pill_{i}"):
                    st.session_state.mid_v_map.pop(i)
                    st.rerun()
        validation_map = st.session_state.mid_v_map

    else:  # Highlighted Values Validation
        detected_cols = set()
        has_highlights = False

        for row in ws1.iter_rows(min_row=2):
            for cell in row:
                if cell.column <= len(headers1):
                    if is_cell_highlighted(cell):
                        col_name = headers1[cell.column - 1]
                        if "Unnamed" not in col_name:
                            detected_cols.add(col_name)
                            has_highlights = True

        if not has_highlights:
            st.warning(
                "⚠️ No highlighted column values detected in this sheet."
            )
        else:
            st.success(f"Detected highlights in: {', '.join(detected_cols)}")
            highlighted_map = [
                {"f1": c, "f2": c, "label": c} for c in detected_cols
            ]

    # Run Comparison
    if st.button("▶ Run Comparison", type="primary", key="run_mid_comp"):
        if not key_f1 or not key_f2:
            st.error("Select Key Columns for both files.")
            return

        final_map = (
            highlighted_map
            if v_mode == "Highlighted Values Validation"
            else validation_map
        )
        if not final_map:
            st.error("No columns selected for validation.")
            return

        t_start = time.perf_counter()

        df2["_k_match"] = df2.apply(
            lambda r: get_row_key(r, key_f2).lower(), axis=1
        )

        results = []
        for idx1, r1 in df1.iterrows():
            excel_row_idx = idx1 + 2
            active_map_for_row = []
            red_map_for_row = []

            for m in final_map:
                if v_mode == "Highlighted Values Validation":
                    try:
                        col_idx = headers1.index(m["f1"]) + 1
                        cell_obj = ws1.cell(row=excel_row_idx, column=col_idx)
                        if is_cell_highlighted(cell_obj):
                            active_map_for_row.append(m)
                            if is_cell_red(cell_obj):
                                red_map_for_row.append(m)
                    except ValueError:
                        pass
                else:
                    active_map_for_row.append(m)

            if v_mode == "Highlighted Values Validation" and not active_map_for_row:
                continue

            k1_match = get_row_key(r1, key_f1).lower()
            k1_display = get_row_key(r1, key_f1)

            matches = df2[df2["_k_match"] == k1_match]
            row_res = {
                "key": k1_display,
                "status": "Match",
                "fields": [],
                "comment": "",
                "is_dup": len(matches) > 1,
                "has_red_fail": False,
            }

            if matches.empty:
                if red_map_for_row:
                    row_res["status"] = "Excluded & Absent"
                else:
                    row_res["status"] = "Missed in Ref File"

                for m in final_map:
                    row_res["fields"].append({
                        "label": m["label"],
                        "v1": (
                            format_val(r1[m["f1"]]) if m["f1"] in r1 else ""
                        ),
                        "v2": "",
                        "match": False,
                        "active": (m in active_map_for_row),
                        "is_red_fail": (m in red_map_for_row),
                    })
            else:
                best_match = None
                best_score = -1

                for _, m_row in matches.iterrows():
                    score = sum(
                        1
                        for m in active_map_for_row
                        if m["f1"] in r1
                        and m["f2"] in m_row
                        and format_val(r1[m["f1"]])
                        == format_val(m_row[m["f2"]])
                    )
                    if score > best_score:
                        best_score, best_match = score, m_row

                any_mismatch = False

                for m in final_map:
                    v1 = format_val(r1[m["f1"]]) if m["f1"] in r1 else ""
                    v2 = (
                        format_val(best_match[m["f2"]])
                        if m["f2"] in best_match
                        else ""
                    )

                    is_red_cell = m in red_map_for_row

                    if m not in active_map_for_row:
                        is_match = True
                    else:
                        is_match = v1 == v2
                        if not is_match:
                            any_mismatch = True

                    if is_red_cell:
                        row_res["has_red_fail"] = True

                    row_res["fields"].append({
                        "label": m["label"],
                        "v1": v1,
                        "v2": v2,
                        "match": is_match,
                        "active": (m in active_map_for_row),
                        "is_red_fail": is_red_cell,
                    })

                if row_res["has_red_fail"]:
                    row_res["status"] = "Exclusion Breach"
                    row_res["comment"] = (
                        "Validation Failed: Row contains excluded record that exists in the reference file."
                    )
                elif any_mismatch:
                    row_res["status"] = "Differs"

                if (
                    not row_res["has_red_fail"]
                    and len(active_map_for_row) > 0
                    and best_score < len(active_map_for_row)
                ):
                    for m in active_map_for_row:
                        if m["f1"] in r1:
                            v1 = format_val(r1[m["f1"]])
                            if v1 != format_val(
                                best_match.get(m["f2"], "")
                            ) and any(
                                format_val(other.get(m["f2"], "")) == v1
                                for _, other in matches.iterrows()
                            ):
                                row_res["comment"] = (
                                    f"{m['label']} value matches a different record in the group."
                                )

            results.append(row_res)

        # Excel Generation
        output = BytesIO()
        wb_out = openpyxl.Workbook()
        ws_all = wb_out.active
        ws_all.title = "All Records"
        ws_mis = wb_out.create_sheet("Mismatches")
        ws_dup = wb_out.create_sheet("Duplicates")

        yellow_fill = PatternFill(
            start_color="FFFF00", end_color="FFFF00", fill_type="solid"
        )
        green_fill = PatternFill(
            start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
        )
        red_fill = PatternFill(
            start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
        )
        bold_font = Font(bold=True)

        f1_n = f1.name.split(".")[0]
        f2_n = f2.name.split(".")[0]

        base_h = ["#", "KEY", "Status"]
        field_h = []
        for m in final_map:
            field_h += [
                f"{m['label']} (Source)",
                f"{m['label']} (Reference)",
                "Match",
            ]
        full_h = base_h + field_h + ["Comment"]

        for ws in [ws_all, ws_dup]:
            ws.append(full_h)
            for cell in ws[1]:
                cell.font = bold_font

        ws_mis.append(
            ["KEY", "Column Name", f"{f1_n} Column", f"{f2_n} Column", "Status"]
        )
        for cell in ws_mis[1]:
            cell.font = bold_font

        idx_all, idx_dup = 1, 1

        for res in results:
            field_vals = []
            for f in res["fields"]:
                field_vals += [
                    f["v1"],
                    f["v2"],
                    "Match" if f["match"] else "Mismatch",
                ]

            # ---- 1. Populate and Highlight the 'All Records' sheet ----
            ws_all.append(
                [idx_all, res["key"], res["status"]]
                + field_vals
                + [res["comment"]]
            )
            curr_row = ws_all.max_row
            idx_all += 1

            if res["status"] == "Match":
                for col_idx in range(1, len(full_h) + 1):
                    ws_all.cell(row=curr_row, column=col_idx).fill = green_fill
            elif res["status"] == "Exclusion Breach":
                ws_all.cell(row=curr_row, column=3).fill = red_fill
            elif res["status"] == "Differs":
                ws_all.cell(row=curr_row, column=3).fill = yellow_fill

            if res["status"] in ["Differs", "Exclusion Breach"]:
                for f_idx, f in enumerate(res["fields"]):
                    is_value_mismatch = format_val(f["v1"]) != format_val(
                        f["v2"]
                    )
                    if is_value_mismatch and f.get("active", True):
                        start_col = 4 + (f_idx * 3)
                        for c in range(start_col, start_col + 3):
                            ws_all.cell(row=curr_row, column=c).fill = (
                                yellow_fill
                            )

                for f in res["fields"]:
                    if f.get("active", True):
                        is_value_mismatch = format_val(f["v1"]) != format_val(
                            f["v2"]
                        )
                        if is_value_mismatch:
                            if res["is_dup"]:
                                status_label = "Mismatch (Duplicate)"
                            else:
                                status_label = (
                                    res["comment"]
                                    if res.get("has_red_fail", False)
                                    else "Mismatch"
                                )
                            ws_mis.append([
                                res["key"],
                                f["label"],
                                f["v1"],
                                f["v2"],
                                status_label,
                            ])

            # ---- 2. Populate and Highlight the 'Duplicates' sheet ----
            if res["is_dup"]:
                ws_dup.append(
                    [idx_dup, res["key"], res["status"]]
                    + field_vals
                    + [res["comment"]]
                )
                curr_dup_row = ws_dup.max_row
                idx_dup += 1

                if res["status"] == "Match":
                    for col_idx in range(1, len(full_h) + 1):
                        ws_dup.cell(
                            row=curr_dup_row, column=col_idx
                        ).fill = green_fill
                elif res["status"] == "Exclusion Breach":
                    ws_dup.cell(row=curr_dup_row, column=3).fill = red_fill
                elif res["status"] == "Differs":
                    ws_dup.cell(row=curr_dup_row, column=3).fill = yellow_fill

                if res["status"] in ["Differs", "Exclusion Breach"]:
                    for f_idx, f in enumerate(res["fields"]):
                        is_value_mismatch = format_val(f["v1"]) != format_val(
                            f["v2"]
                        )
                        if is_value_mismatch and f.get("active", True):
                            start_col = 4 + (f_idx * 3)
                            for c in range(start_col, start_col + 3):
                                ws_dup.cell(
                                    row=curr_dup_row, column=c
                                ).fill = yellow_fill

        wb_out.save(output)

        elapsed = time.perf_counter() - t_start
        mismatches_count = ws_mis.max_row - 1 if ws_mis.max_row > 1 else 0

        log_event(
            feature="MID File Validator",
            duration_sec=elapsed,
            details={
                "source_file": f1.name,
                "target_file": f2.name,
                "source_rows": len(df1),
                "target_rows": len(df2),
                "mismatches": mismatches_count,
            },
        )

        st.success(f"Validation finished in {elapsed:.2f}s.")
        st.download_button(
            "⬇ Download Results",
            data=output.getvalue(),
            file_name="Validation_Results.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


if __name__ == "__main__":
    run()