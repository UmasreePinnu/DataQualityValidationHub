import streamlit as st
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def run():
    st.title("📄 Validation Report Generator")

    # 1. File Configuration
    file_name_input = st.text_input("Enter Output File Name",
                                    value=f"Validation_Report_{datetime.now().strftime('%Y%m%d')}")

    # 2. Scenario Selection
    st.subheader("Select Scenarios")
    available_scenarios = [
        "Table Check",
        "Column Check",
        "Datatype Check",
        "Information Schema VS Column Mappings Check",
        "Case Sensitive Check",
        "Row Count Check",
        "Row To Row Validations",
        "Duplicate Check",
        "Null Count Check",
        "Fact Metric Validations",
        "Keys Validations"
    ]

    selected_scenarios = []
    cols = st.columns(3)
    for i, s in enumerate(available_scenarios):
        if cols[i % 3].checkbox(s, key=f"scen_check_{i}"):
            selected_scenarios.append(s)

    # 3. Generation Logic
    if st.button("🚀 Generate Excel Report"):
        if not selected_scenarios:
            st.warning("Please select at least one scenario.")
            return

        output = BytesIO()
        wb = Workbook()

        # --- SUMMARY SHEET ---
        ws_summary = wb.active
        ws_summary.title = "Summary"

        headers = ["Scenario Name", "Status", "Dev Comments", "QA Comments"]
        ws_summary.append(headers)

        # Styling: Bold, Thick Dark Blue Cell Color, White Font
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                             top=Side(style='thin'), bottom=Side(style='thin'))

        for col_num, header in enumerate(headers, 1):
            cell = ws_summary.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border
            ws_summary.column_dimensions[cell.column_letter].width = 25

        # Fill Summary Rows
        for s_name in selected_scenarios:
            ws_summary.append([s_name, "", "", ""])
            # Apply borders to the new row
            for col in range(1, 5):
                ws_summary.cell(row=ws_summary.max_row, column=col).border = thin_border

        # --- SEPARATE SCENARIO SHEETS ---
        for s_name in selected_scenarios:
            # Excel limits sheet names to 31 characters
            wb.create_sheet(title=s_name[:30])

        wb.save(output)
        st.success(f"Successfully generated {file_name_input}.xlsx")
        st.download_button(
            label="📥 Download Validation Report",
            data=output.getvalue(),
            file_name=f"{file_name_input}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )